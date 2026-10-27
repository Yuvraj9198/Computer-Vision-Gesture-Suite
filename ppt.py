print("ppt.py is running")

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor


def create_presentation():
    prs = Presentation()

    # Helper function to add slides
    def add_slide(title_text, content_points):
        slide_layout = prs.slide_layouts[1]  # Title and Content
        slide = prs.slides.add_slide(slide_layout)

        # Title
        title = slide.shapes.title
        title.text = title_text
        title_tf = title.text_frame
        title_tf.paragraphs[0].font.size = Pt(36)
        title_tf.paragraphs[0].font.bold = True
        title_tf.paragraphs[0].font.color.rgb = RGBColor(0, 51, 102)

        # Content
        body_tf = slide.shapes.placeholders[1].text_frame
        body_tf.clear()

        for i, point in enumerate(content_points):
            if i == 0:
                p = body_tf.paragraphs[0]
            else:
                p = body_tf.add_paragraph()

            p.text = point
            p.font.size = Pt(24)

    # SLIDE 1: TITLE SLIDE
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "IoT Pathshala"
    slide.placeholders[1].text = (
        "Hands-On Learning for the Future\n"
        "Empowering Students Through Practical Innovation"
    )

    # SLIDE 2: WHY IOT PATHSHALA
    add_slide(
        "Why IoT Pathshala?",
        [
            "Transforming Curiosity into Capability",
            "Students need more than theory — they need hands-on exposure",
            "Bridging the gap between textbooks and real-world technology",
            "Aligned with Atal Tinkering Lab (ATL) objectives",
            "Building strong STEM and problem-solving foundations",
        ],
    )

    # SLIDE 3: WHAT WE OFFER
    add_slide(
        "Our Core Offerings",
        [
            "Hands-On IoT Workshops (Grades 6–12)",
            "Beginner-friendly — no prior coding required",
            "Project-based learning with sensors and automation",
            "ATL Lab collaboration and mentorship",
            "Guidance for student-led innovation projects",
        ],
    )

    # SLIDE 4: STUDENT OUTCOMES
    add_slide(
        "Student Outcomes",
        [
            "Students will be able to:",
            "Understand sensors and microcontrollers",
            "Build simple automation systems",
            "Read and interpret real sensor data",
            "Develop confidence in experimenting with technology",
        ],
    )

    # SLIDE 5: STRUCTURE & SAFETY
    add_slide(
        "Workshop Structure & Safety",
        [
            "Flexible Workshop Formats",
            "Half-day introductory session",
            "1-day intensive hands-on program",
            "3–5 day project-based module",
            "Age-appropriate components with guided supervision",
        ],
    )

    # SLIDE 6: CONTACT (PPT 6)
    add_slide(
        "Let’s Partner for Impact",
        [
            "Shape curious minds into confident innovators",
            "Contact Us",
            "IoT Pathshala",
            "www.iotpathshala.com",
            "Phone: [Insert Phone Number]",
            "Email: [Insert Email Address]",
        ],
    )

    # SAVE FILE
    file_name = "IoT_Pathshala_Pitch.pptx"
    prs.save(file_name)
    print(f"Presentation saved as {file_name}")


if __name__ == "__main__":
    create_presentation()
